"""格式规则的阶段适用性与「必需章节」分阶段默认值测试。"""
from __future__ import annotations

from pathlib import Path

from docx import Document

from tools.docx_parser import parse_docx
from tools.format_expert import FormatExpert
from tools.rules_schema import (
    FINAL_ONLY_RULES,
    RULES,
    STAGE_SECTION_PRESETS,
    rule_applies_to_stage,
    stage_sections_expected,
)


def _rules() -> dict:
    """与 competition.aip_partner._default_rules 等价的默认规则配置。"""
    return {
        rule.key: {
            "category": rule.category,
            "title": rule.title,
            "enabled": bool(rule.preset),
            "expected": rule.preset,
            "weight": float(rule.default_weight),
        }
        for rule in RULES
    }


def _active_count(rules: dict, stage: str | None) -> int:
    return sum(
        1
        for key, cfg in rules.items()
        if cfg["enabled"] and cfg["expected"] and rule_applies_to_stage(key, stage)
    )


def _make_proposal_doc(path: Path) -> None:
    doc = Document()
    doc.add_heading("一、研究背景", level=1)
    doc.add_paragraph("随着毕业论文管理信息化推进，质量检查需要自动化手段支撑。" * 8)
    doc.add_heading("二、研究内容", level=1)
    doc.add_paragraph("本文设计并实现论文格式与逻辑检查系统。" * 8)
    doc.save(path)


def test_rule_applies_to_stage_mapping() -> None:
    assert rule_applies_to_stage("toc_required", "proposal") is False
    assert rule_applies_to_stage("abstract_word_min", "midterm") is False
    assert rule_applies_to_stage("ref_min_count", "proposal") is False
    assert rule_applies_to_stage("toc_required", "final") is True
    assert rule_applies_to_stage("body_font_cn", "proposal") is True
    assert rule_applies_to_stage("body_font_cn", "midterm") is True
    assert rule_applies_to_stage("ref_min_count", None) is True
    assert rule_applies_to_stage("required_sections", "proposal") is True


def test_stage_sections_expected_substitution() -> None:
    default = next(r for r in RULES if r.key == "required_sections").preset
    assert stage_sections_expected("proposal", default) == STAGE_SECTION_PRESETS["proposal"]
    assert stage_sections_expected("midterm", default) == STAGE_SECTION_PRESETS["midterm"]
    assert stage_sections_expected("final", default) == default
    assert stage_sections_expected(None, default) == default
    # 教师自定义的章节清单不被替换
    custom = "摘要,目录,结论"
    assert stage_sections_expected("proposal", custom) == custom


def test_proposal_stage_skips_final_only_rules(tmp_path: Path) -> None:
    path = tmp_path / "proposal.docx"
    _make_proposal_doc(path)
    rules = _rules()
    report = FormatExpert(rules).check(parse_docx(path), stage="proposal")

    assert report.stage == "proposal"
    by_key = {item.key: item for item in report.items}
    # 仅终稿适用的规则被标记为不适用且状态为 na
    for key in FINAL_ONLY_RULES & set(by_key):
        assert by_key[key].applicable is False
        assert by_key[key].status == "na"
    # 不适用规则不产生问题，也不计入统计
    assert not [i for i in report.issues if i.rule_key in FINAL_ONLY_RULES]
    assert report.configured_count == _active_count(rules, "proposal")
    assert report.configured_count < _active_count(rules, "final")
    # 「必需章节」按开题阶段默认值替换
    assert by_key["required_sections"].expected == STAGE_SECTION_PRESETS["proposal"]
    assert by_key["required_sections"].applicable is True
    assert report.score is not None


def test_final_stage_counts_everything(tmp_path: Path) -> None:
    path = tmp_path / "final.docx"
    _make_proposal_doc(path)
    rules = _rules()
    report = FormatExpert(rules).check(parse_docx(path), stage="final")

    assert all(item.applicable for item in report.items)
    assert report.configured_count == _active_count(rules, "final")
    # 终稿默认值不替换
    by_key = {item.key: item for item in report.items}
    default = next(r for r in RULES if r.key == "required_sections").preset
    assert by_key["required_sections"].expected == default


def test_no_stage_is_backward_compatible(tmp_path: Path) -> None:
    path = tmp_path / "plain.docx"
    _make_proposal_doc(path)
    rules = _rules()
    report = FormatExpert(rules).check(parse_docx(path))

    assert report.stage is None
    assert all(item.applicable for item in report.items)
    assert report.configured_count == _active_count(rules, None)


def test_report_html_marks_na_rules(tmp_path: Path) -> None:
    """报告页面对不适用规则给出提示，且不计入图表数据。"""
    from tools.logic_expert import LogicExpert
    from tools.models import AnalysisResult
    from tools.report_builder import ReportBuilder

    path = tmp_path / "proposal.docx"
    _make_proposal_doc(path)
    doc = parse_docx(path)
    rules = _rules()
    fmt = FormatExpert(rules).check(doc, stage="proposal")
    logic = LogicExpert().analyze(doc, stage="proposal")
    result = AnalysisResult(
        paper_id=1, paper_title="样例论文", student_name="学生甲",
        format_report=fmt, logic_report=logic,
    )
    template_dir = Path(__file__).resolve().parent.parent / "templates"
    builder = ReportBuilder(template_dir, tmp_path / "rep", tmp_path / "json")
    html_path, _ = builder.build(result, "演示高校")
    html = Path(html_path).read_text(encoding="utf-8")

    assert "默认不适用" in html
    assert "开题报告" in html
    # 不适用规则的标题出现在提示中，但不应出现在图表标签 JSON 中
    assert "目录须自动生成" in html
    assert html.count("目录须自动生成") == 1


def test_report_html_final_stage_has_no_na_note(tmp_path: Path) -> None:
    """终稿阶段所有规则适用，报告不出现不适用提示。"""
    from tools.logic_expert import LogicExpert
    from tools.models import AnalysisResult
    from tools.report_builder import ReportBuilder

    path = tmp_path / "final.docx"
    _make_proposal_doc(path)
    doc = parse_docx(path)
    fmt = FormatExpert(_rules()).check(doc, stage="final")
    logic = LogicExpert().analyze(doc, stage="final")
    result = AnalysisResult(
        paper_id=2, paper_title="样例论文", student_name="学生乙",
        format_report=fmt, logic_report=logic,
    )
    template_dir = Path(__file__).resolve().parent.parent / "templates"
    builder = ReportBuilder(template_dir, tmp_path / "rep", tmp_path / "json")
    html_path, _ = builder.build(result, "演示高校")
    html = Path(html_path).read_text(encoding="utf-8")
    assert "默认不适用" not in html
