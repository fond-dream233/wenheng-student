from __future__ import annotations

import io
import re
import uuid
from unittest.mock import patch

from docx import Document

from app import app
from tools.database import now_str
from utils.security import hash_password


def _csrf_token(html: str) -> str:
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert match
    return match.group(1)


def _docx_bytes(topic: str) -> io.BytesIO:
    stream = io.BytesIO()
    document = Document()
    document.add_heading(topic, level=0)
    document.add_heading('研究目标', level=1)
    document.add_paragraph((f'围绕{topic}建立可解释的研究目标与评价指标。') * 20)
    document.add_heading('研究方法', level=1)
    document.add_paragraph((f'采用文档解析和规则匹配方法研究{topic}。') * 20)
    document.add_heading('结论', level=1)
    document.add_paragraph((f'{topic}实验已经完成并形成可复核结果。') * 15)
    document.save(stream)
    stream.seek(0)
    return stream


def _student(client, username: str) -> int:
    db = app.extensions['db']
    user_id = db.execute(
        'INSERT INTO users (username, password_hash, role, name, student_no, created_at)'
        " VALUES (?, ?, 'student', ?, ?, ?)",
        (username, hash_password('Test-password-123'), '匿名学生', username, now_str()),
    )
    with client.session_transaction() as session:
        session['user_id'] = user_id
        session['username'] = username
        session['role'] = 'student'
        session['name'] = '匿名学生'
    return user_id


def test_student_uploads_stages_and_views_cross_stage_report():
    client = app.test_client()
    username = f"stage-{uuid.uuid4().hex[:10]}"
    student_id = _student(client, username)
    token = _csrf_token(client.get('/student/').get_data(as_text=True))

    with patch('routes.student.Analyzer.run'):
        proposal = client.post('/student/upload', data={
            'csrf_token': token,
            'project_title': '匿名课题 A',
            'stage': 'proposal',
            'file': (_docx_bytes('论文质量检查'), 'proposal.docx'),
        }, content_type='multipart/form-data')
        final = client.post('/student/upload', data={
            'csrf_token': token,
            'project_title': '匿名课题 A',
            'stage': 'final',
            'file': (_docx_bytes('论文质量检查与多智能体协作'), 'final.docx'),
        }, content_type='multipart/form-data')
    assert proposal.status_code == 302
    assert final.status_code == 302

    db = app.extensions['db']
    project = db.get(
        'SELECT * FROM thesis_projects WHERE student_id = ? AND title = ?',
        (student_id, '匿名课题 A'),
    )
    assert project
    papers = db.query('SELECT stage FROM papers WHERE project_id = ? ORDER BY id', (project['id'],))
    assert [row['stage'] for row in papers] == ['proposal', 'final']

    response = client.post(
        f"/student/projects/{project['id']}/compare",
        data={'csrf_token': token},
        follow_redirects=True,
    )
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert '跨阶段分析报告' in html
    assert '开题' in html and '终稿' in html
    assert '综合漂移分' in html
    report = db.get('SELECT * FROM stage_reports WHERE project_id = ?', (project['id'],))
    assert report and 0 <= report['drift_score'] <= 100

    other_client = app.test_client()
    _student(other_client, f"other-{uuid.uuid4().hex[:10]}")
    assert other_client.get(f"/student/stage-report/{report['id']}").status_code == 404

    teacher_client = app.test_client()
    with teacher_client.session_transaction() as session:
        session['user_id'] = 1
        session['role'] = 'teacher'
        session['name'] = '教师'
    teacher_page = teacher_client.get(f"/teacher/stage-report/{report['id']}")
    assert teacher_page.status_code == 200
    assert '匿名课题 A' in teacher_page.get_data(as_text=True)


def test_compare_requires_two_distinct_stages():
    client = app.test_client()
    username = f"single-{uuid.uuid4().hex[:10]}"
    student_id = _student(client, username)
    db = app.extensions['db']
    project_id = db.execute(
        'INSERT INTO thesis_projects (student_id, title, created_at, updated_at)'
        ' VALUES (?, ?, ?, ?)',
        (student_id, '单阶段课题', now_str(), now_str()),
    )
    token = _csrf_token(client.get('/student/').get_data(as_text=True))
    response = client.post(
        f'/student/projects/{project_id}/compare',
        data={'csrf_token': token},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert '至少上传两个不同阶段' in response.get_data(as_text=True)
