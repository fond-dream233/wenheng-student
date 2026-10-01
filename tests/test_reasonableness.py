"""内容合理性检查与跨阶段引用的集成测试。"""
from __future__ import annotations

from pathlib import Path

from docx import Document

from competition.stage_comparison import compare_stage_documents
from tools.docx_parser import parse_docx
from tools.logic_expert import LogicExpert


def _build(path: Path, sections: list[tuple[str, str]]) -> None:
    doc = Document()
    for title, body in sections:
        doc.add_heading(title, level=1)
        doc.add_paragraph(body)
    doc.save(path)


def test_proposal_objective_specificity(tmp_path: Path) -> None:
    path = tmp_path / "proposal.docx"
    _build(path, [
        ("第一章 绪论", "本文拟开展相关研究工作。" * 20),
        ("研究目标", "本文希望对相关问题进行较为深入的研究，力争取得较好的成果。" * 3),
    ])
    report = LogicExpert().analyze(parse_docx(path), stage="proposal")
    keys = {f.key for f in report.findings}
    assert "objective_specificity" in keys
    finding = next(f for f in report.findings if f.key == "objective_specificity")
    assert finding.category == "reasonableness"
    assert finding.citation is not None


def test_conclusion_echo_mismatch(tmp_path: Path) -> None:
    path = tmp_path / "final.docx"
    _build(path, [
        ("研究目标", "本文研究论文质量检查方法，构建规则校验与格式审查体系。" * 10),
        ("结论", "量子拓扑绝缘体材料在低温下表现出优异的输运特性与霍尔效应。" * 10),
    ])
    report = LogicExpert().analyze(parse_docx(path), stage="final")
    finding = next((f for f in report.findings if f.key == "conclusion_echo"), None)
    assert finding is not None
    assert finding.category == "reasonableness"
    assert finding.metrics.get("conclusion_objective_sim", 1.0) < 0.35


def test_midterm_result_alignment(tmp_path: Path) -> None:
    path = tmp_path / "midterm.docx"
    _build(path, [
        ("研究目标", "本文研究基于规则的论文格式自动检查方法。" * 10),
        ("阶段性成果", "已完成为期两周的野外地质采样与岩石薄片显微鉴定工作。" * 10),
    ])
    report = LogicExpert().analyze(parse_docx(path), stage="midterm")
    finding = next((f for f in report.findings if f.key == "result_objective_alignment"), None)
    assert finding is not None
    assert finding.category == "reasonableness"


def test_reasonableness_skips_other_stages(tmp_path: Path) -> None:
    path = tmp_path / "proposal.docx"
    _build(path, [
        ("第一章 绪论", "本文研究论文质量检查方法。" * 30),
        ("研究目标", "建立可衡量的检查指标并通过实验验证准确率。" * 10),
        ("结论", "实验表明系统有效。" * 10),
    ])
    report = LogicExpert().analyze(parse_docx(path), stage="proposal")
    assert "conclusion_echo" not in {f.key for f in report.findings}


def test_cross_stage_findings_carry_citation(tmp_path: Path) -> None:
    proposal = tmp_path / "proposal.docx"
    final = tmp_path / "final.docx"
    _build(proposal, [("第一章 绪论", "本文研究基于规则的本科论文质量检查方法与技术路线。" * 20)])
    _build(final, [("第一章 绪论", "本研究探讨热带海洋珊瑚礁生态系统的生物多样性保护策略。" * 20)])
    result = compare_stage_documents([
        ("proposal", proposal, "proposal.docx"),
        ("final", final, "final.docx"),
    ])
    findings = result["report"]["findings"]
    assert findings, "主题完全不同的两阶段文档应产生漂移发现项"
    assert all("citation" in f and f["citation"] for f in findings)
