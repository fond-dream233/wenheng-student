from __future__ import annotations

import csv
from pathlib import Path

import pytest
from docx import Document

from evaluation.blind_review import aggregate_reviews, prepare_package


def _write_doc(path: Path, identity: str, topic: str) -> None:
    document = Document()
    document.core_properties.author = identity
    document.add_heading(f"{identity}的{topic}", level=0)
    document.add_heading("研究目标", level=1)
    document.add_paragraph(f"{identity}研究{topic}的连续发展。")
    document.save(path)


def _write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def test_prepare_package_separates_mapping_and_redacts(tmp_path: Path):
    source = tmp_path / "source"
    source.mkdir()
    _write_doc(source / "proposal.docx", "张三", "智能灌溉")
    _write_doc(source / "final.docx", "张三", "智能灌溉")
    manifest = source / "manifest.csv"
    _write_csv(manifest, [{
        "source_case_id": "student-001",
        "proposal_path": "proposal.docx",
        "midterm_path": "",
        "final_path": "final.docx",
        "sensitive_terms": "张三|某大学",
        "consent_confirmed": "yes",
    }])

    public = tmp_path / "public"
    private = tmp_path / "private"
    receipt = prepare_package(manifest, public, private, seed=42)

    assert receipt["caseCount"] == 1
    assert not (public / "identity_mapping.csv").exists()
    assert "student-001" in (private / "identity_mapping.csv").read_text(encoding="utf-8-sig")
    redacted = Document(public / "documents" / "BR001" / "proposal.docx")
    assert "张三" not in "\n".join(paragraph.text for paragraph in redacted.paragraphs)
    assert redacted.core_properties.author == "Anonymous Review"


def _review_row(reviewer: str, case_id: str, label: str) -> dict[str, str]:
    return {
        "reviewer_code": reviewer,
        "case_id": case_id,
        "available_stages": "proposal|final",
        "topic_continuity_1_5": "4",
        "objective_continuity_1_5": "4",
        "method_continuity_1_5": "3",
        "scope_change": "refinement" if label == "low" else "partial",
        "overall_risk": label,
        "confidence_1_5": "4",
        "evidence_location": "研究目标章节",
        "rationale": "研究主线延续，但方法有调整。",
        "review_complete": "yes",
    }


def test_aggregate_requires_adjudication_for_conflicts(tmp_path: Path):
    reviewer_a = tmp_path / "a.csv"
    reviewer_b = tmp_path / "b.csv"
    _write_csv(reviewer_a, [_review_row("A", "BR001", "low"), _review_row("A", "BR002", "medium")])
    _write_csv(reviewer_b, [_review_row("B", "BR001", "low"), _review_row("B", "BR002", "high")])

    output = tmp_path / "aggregate"
    report = aggregate_reviews(reviewer_a, reviewer_b, output)
    assert report["rawAgreement"] == 0.5
    assert report["conflictCount"] == 1
    assert report["baselineComplete"] is False

    adjudication = output / "adjudication.csv"
    rows = list(csv.DictReader(adjudication.open(encoding="utf-8-sig")))
    rows[0]["final_risk"] = "medium"
    rows[0]["final_reason"] = "保留主线但核心目标部分变化。"
    rows[0]["adjudicator"] = "C"
    _write_csv(adjudication, rows)
    final = aggregate_reviews(reviewer_a, reviewer_b, tmp_path / "final", adjudication_path=adjudication)
    assert final["baselineComplete"] is True
    assert final["finalizedCount"] == 2


def test_prepare_rejects_unconfirmed_consent(tmp_path: Path):
    manifest = tmp_path / "manifest.csv"
    _write_csv(manifest, [{
        "source_case_id": "student-001",
        "proposal_path": "missing.docx",
        "midterm_path": "",
        "final_path": "missing.docx",
        "sensitive_terms": "",
        "consent_confirmed": "no",
    }])
    with pytest.raises(ValueError, match="未确认"):
        prepare_package(manifest, tmp_path / "out", tmp_path / "private")


def test_aggregate_rejects_same_reviewer_code(tmp_path: Path):
    reviewer_a = tmp_path / "a.csv"
    reviewer_b = tmp_path / "b.csv"
    _write_csv(reviewer_a, [_review_row("A", "BR001", "low")])
    _write_csv(reviewer_b, [_review_row("A", "BR001", "low")])
    with pytest.raises(ValueError, match="不同"):
        aggregate_reviews(reviewer_a, reviewer_b, tmp_path / "aggregate")
