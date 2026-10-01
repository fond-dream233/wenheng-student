from __future__ import annotations

import base64
import json
import uuid
from datetime import datetime, timezone
from pathlib import Path

import pytest
from acps_sdk.aip import StructuredDataItem, TaskCommand, TaskCommandType, TaskState, TextDataItem
from acps_sdk.aip.aip_rpc_server import TaskManager
from docx import Document

from competition.aip_partner import build_handlers


@pytest.fixture(autouse=True)
def clear_tasks():
    TaskManager._tasks.clear()
    yield
    TaskManager._tasks.clear()


@pytest.fixture()
def docx_payload(tmp_path: Path) -> str:
    path = tmp_path / "paper.docx"
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("本文研究论文质量检查方法，并给出实验数据和参考文献支撑。" * 30)
    doc.add_heading("结论", level=1)
    doc.add_paragraph("实验结果表明系统可以输出可定位的修改建议。" * 20)
    doc.save(path)
    return json.dumps({"filename": path.name,
                       "fileBase64": base64.b64encode(path.read_bytes()).decode("ascii")})


def _stage_document(path: Path, topic: str, objective: str, method: str, conclusion: str) -> str:
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph(topic * 20)
    doc.add_heading("研究目标", level=1)
    doc.add_paragraph(objective * 15)
    doc.add_heading("研究方法", level=1)
    doc.add_paragraph(method * 15)
    doc.add_heading("结论", level=1)
    doc.add_paragraph(conclusion * 10)
    doc.save(path)
    return base64.b64encode(path.read_bytes()).decode("ascii")


@pytest.fixture()
def stage_payload(tmp_path: Path) -> str:
    proposal = tmp_path / "proposal.docx"
    final = tmp_path / "final.docx"
    shared_topic = "本文研究基于规则与多智能体协作的本科论文质量检查方法。"
    return json.dumps({
        "documents": [
            {
                "stage": "proposal",
                "filename": proposal.name,
                "fileBase64": _stage_document(
                    proposal,
                    shared_topic,
                    "建立可解释的格式检查与逻辑审查体系。",
                    "采用文档解析、规则匹配和多智能体任务协作。",
                    "开题阶段形成研究计划。",
                ),
            },
            {
                "stage": "final",
                "filename": final.name,
                "fileBase64": _stage_document(
                    final,
                    shared_topic + "系统已经完成实现与测试。",
                    "建立可解释的格式检查与逻辑审查体系并完成验证。",
                    "采用文档解析、规则匹配和多智能体任务协作。",
                    "实验表明系统能够输出可定位的检查建议。",
                ),
            },
        ]
    }, ensure_ascii=False)


def command(kind: TaskCommandType, task_id: str, text: str | None = None) -> TaskCommand:
    return TaskCommand(
        id=f"cmd-{uuid.uuid4()}",
        sentAt=datetime.now(timezone.utc).isoformat(),
        senderRole="leader",
        senderId="test-leader",
        sessionId="test-session",
        taskId=task_id,
        command=kind,
        dataItems=[TextDataItem(text=text)] if text is not None else None,
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("role", ["format", "logic"])
async def test_start_get_complete_protocol(role: str, docx_payload: str):
    handlers = build_handlers(role, f"test-{role}-aic")
    task_id = f"task-{role}"
    started = await handlers.on_start(command(TaskCommandType.Start, task_id, docx_payload), None)
    assert started.status.state == TaskState.AwaitingCompletion
    assert started.senderId == f"test-{role}-aic"
    assert started.products and started.products[0].dataItems

    fetched = await handlers.on_get(command(TaskCommandType.Get, task_id), started)
    assert fetched.status.state == TaskState.AwaitingCompletion
    completed = await handlers.on_complete(command(TaskCommandType.Complete, task_id), fetched)
    assert completed.status.state == TaskState.Completed


@pytest.mark.asyncio
async def test_awaiting_input_continue_and_cancel(docx_payload: str):
    handlers = build_handlers("format", "test-format-aic")
    waiting = await handlers.on_start(command(TaskCommandType.Start, "needs-input", "{}"), None)
    assert waiting.status.state == TaskState.AwaitingInput
    continued = await handlers.on_continue(
        command(TaskCommandType.Continue, "needs-input", docx_payload), waiting
    )
    assert continued.status.state == TaskState.AwaitingCompletion

    other = await handlers.on_start(command(TaskCommandType.Start, "cancel-me", "{}"), None)
    canceled = await handlers.on_cancel(command(TaskCommandType.Cancel, "cancel-me"), other)
    assert canceled.status.state == TaskState.Canceled


@pytest.mark.asyncio
async def test_rejects_untrusted_local_path(tmp_path: Path):
    handlers = build_handlers("logic", "test-logic-aic")
    payload = json.dumps({"documentPath": str(tmp_path / "paper.docx")})
    failed = await handlers.on_start(command(TaskCommandType.Start, "unsafe-path", payload), None)
    assert failed.status.state == TaskState.Failed
    assert "本地路径输入已禁用" in failed.status.dataItems[0].text


@pytest.mark.asyncio
async def test_cross_stage_partner_returns_explainable_report(stage_payload: str):
    handlers = build_handlers("cross_stage", "test-cross-stage-aic")
    started = await handlers.on_start(
        command(TaskCommandType.Start, "cross-stage-task", stage_payload), None
    )
    assert started.status.state == TaskState.AwaitingCompletion
    item = started.products[0].dataItems[0]
    assert isinstance(item, StructuredDataItem)
    result = item.data
    assert result["agent"] == "cross_stage"
    assert result["report"]["stageCount"] == 2
    assert len(result["report"]["comparisons"]) == 1
    assert 0 <= result["report"]["overallDriftScore"] <= 100
    assert result["report"]["riskLevel"] in {"low", "medium", "high"}
    assert "不等同于语义事实判断" in result["report"]["disclaimer"]

    completed = await handlers.on_complete(
        command(TaskCommandType.Complete, "cross-stage-task"), started
    )
    assert completed.status.state == TaskState.Completed


@pytest.mark.asyncio
async def test_cross_stage_rejects_duplicate_stage(stage_payload: str):
    payload = json.loads(stage_payload)
    payload["documents"][1]["stage"] = "proposal"
    handlers = build_handlers("cross_stage", "test-cross-stage-aic")
    failed = await handlers.on_start(
        command(TaskCommandType.Start, "duplicate-stage", json.dumps(payload)), None
    )
    assert failed.status.state == TaskState.Failed
    assert "阶段 proposal 重复" in failed.status.dataItems[0].text
