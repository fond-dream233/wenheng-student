"""知识库引用与阶段要素检查的集成测试。"""
from __future__ import annotations

from pathlib import Path

from docx import Document

from tools.docx_parser import parse_docx
from tools.knowledge_base import cite_for
from tools.logic_expert import LogicExpert


def _make_proposal(path: Path) -> None:
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("本文研究论文质量检查方法。" * 30)
    doc.add_heading("结论", level=1)
    doc.add_paragraph("实验表明系统可以输出修改建议。" * 20)
    doc.save(path)


def test_cite_for_returns_stage_clause() -> None:
    cited = cite_for("开题报告 研究目标 技术路线", stage="proposal")
    assert cited is not None
    assert cited["id"].startswith("proposal")
    assert cited["clause"].strip()


def test_logic_findings_carry_citation(tmp_path: Path) -> None:
    path = tmp_path / "proposal.docx"
    _make_proposal(path)
    report = LogicExpert().analyze(parse_docx(path), stage="proposal")

    stage_findings = [f for f in report.findings if f.key == "stage_requirements"]
    assert stage_findings, "proposal 阶段缺少研究目标/技术路线时应产生要素缺失项"
    citation = stage_findings[0].citation
    assert citation is not None
    assert citation["clause"].strip()
    assert citation["stage"] in {"proposal", "common"}


def test_stage_requirements_key_present(tmp_path: Path) -> None:
    path = tmp_path / "proposal.docx"
    _make_proposal(path)
    report = LogicExpert().analyze(parse_docx(path), stage="proposal")
    assert "stage_requirements" in {finding.key for finding in report.findings}
