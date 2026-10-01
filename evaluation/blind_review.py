"""Prepare and aggregate independent teacher blind reviews of real stage documents.

The workflow is deliberately local-only. It rebuilds clean DOCX files from visible
text, keeps the identity mapping outside the review package, and never runs the
cross-stage model before the human labels are locked.
"""
from __future__ import annotations

import argparse
import csv
import json
import random
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from docx import Document

STAGES = ("proposal", "midterm", "final")
RISK_LABELS = ("low", "medium", "high")
REVIEW_LABELS = (*RISK_LABELS, "uncertain")
REVIEW_COLUMNS = (
    "reviewer_code",
    "case_id",
    "available_stages",
    "topic_continuity_1_5",
    "objective_continuity_1_5",
    "method_continuity_1_5",
    "scope_change",
    "overall_risk",
    "confidence_1_5",
    "evidence_location",
    "rationale",
    "review_complete",
)
MANIFEST_COLUMNS = (
    "source_case_id",
    "proposal_path",
    "midterm_path",
    "final_path",
    "sensitive_terms",
    "consent_confirmed",
)
SCOPE_VALUES = {"none", "refinement", "partial", "replacement", "uncertain"}
YES_VALUES = {"1", "true", "yes", "y", "是", "已确认"}


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _write_csv(path: Path, rows: Iterable[dict[str, Any]], columns: Iterable[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(columns), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _redact(text: str, terms: Iterable[str]) -> str:
    value = text
    for term in sorted({item.strip() for item in terms if item.strip()}, key=len, reverse=True):
        value = re.sub(re.escape(term), "[已匿名]", value, flags=re.IGNORECASE)
    return value


def _visible_paragraphs(document: Any) -> Iterable[Any]:
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _build_clean_docx(source: Path, target: Path, sensitive_terms: list[str]) -> None:
    """Rebuild a text-only DOCX so comments, images and hidden metadata cannot leak."""
    original = Document(source)
    clean = Document()
    if clean.paragraphs and not clean.paragraphs[0].text:
        paragraph = clean.paragraphs[0]
        first = True
    else:
        paragraph = None
        first = False

    for item in original.paragraphs:
        text = _redact(item.text, sensitive_terms)
        output = paragraph if first and paragraph is not None else clean.add_paragraph()
        first = False
        output.text = text
        try:
            output.style = item.style.name
        except (AttributeError, KeyError):
            pass

    for table in original.tables:
        clean.add_paragraph("[表格内容]")
        for row in table.rows:
            values = [_redact(cell.text, sensitive_terms).strip() for cell in row.cells]
            clean.add_paragraph(" | ".join(values))

    properties = clean.core_properties
    properties.author = "Anonymous Review"
    properties.last_modified_by = "Anonymous Review"
    properties.title = ""
    properties.subject = ""
    properties.keywords = ""
    properties.comments = ""
    properties.category = ""
    target.parent.mkdir(parents=True, exist_ok=True)
    clean.save(target)

    reopened = Document(target)
    visible = "\n".join(paragraph.text for paragraph in _visible_paragraphs(reopened)).lower()
    remaining = [term for term in sensitive_terms if term.strip().lower() in visible]
    if remaining:
        target.unlink(missing_ok=True)
        raise ValueError(f"匿名化失败，文档仍含敏感词：{', '.join(remaining)}")


def _validate_source_manifest(rows: list[dict[str, str]], manifest_path: Path) -> None:
    if not rows:
        raise ValueError("源清单没有样本")
    missing_columns = set(MANIFEST_COLUMNS) - set(rows[0])
    if missing_columns:
        raise ValueError(f"源清单缺少列：{', '.join(sorted(missing_columns))}")
    ids = [row["source_case_id"].strip() for row in rows]
    if any(not value for value in ids) or len(ids) != len(set(ids)):
        raise ValueError("source_case_id 必须非空且唯一")
    for row in rows:
        if row["consent_confirmed"].strip().lower() not in YES_VALUES:
            raise ValueError(f"样本 {row['source_case_id']} 未确认授权/知情同意")
        stage_count = 0
        for stage in STAGES:
            raw_path = row[f"{stage}_path"].strip()
            if not raw_path:
                continue
            stage_count += 1
            source = Path(raw_path)
            if not source.is_absolute():
                source = manifest_path.parent / source
            if source.suffix.lower() != ".docx" or not source.is_file():
                raise ValueError(f"样本 {row['source_case_id']} 的 {stage} 文件无效：{source}")
        if stage_count < 2:
            raise ValueError(f"样本 {row['source_case_id']} 至少需要两个阶段")


def prepare_package(
    manifest_path: Path,
    output_dir: Path,
    private_dir: Path,
    *,
    seed: int | None = None,
) -> dict[str, Any]:
    rows = _read_csv(manifest_path)
    _validate_source_manifest(rows, manifest_path)
    resolved_output = output_dir.resolve()
    resolved_private = private_dir.resolve()
    if (resolved_output == resolved_private
            or resolved_private.is_relative_to(resolved_output)
            or resolved_output.is_relative_to(resolved_private)):
        raise ValueError("盲审包与私密映射目录必须分开且不能相互嵌套")
    if any(output_dir.iterdir()) if output_dir.exists() else False:
        raise ValueError(f"输出目录必须为空：{output_dir}")

    actual_seed = seed if seed is not None else random.SystemRandom().randrange(1, 2**31)
    assignment = list(rows)
    random.Random(actual_seed).shuffle(assignment)
    case_manifest: list[dict[str, str]] = []
    mapping: list[dict[str, str]] = []
    for index, row in enumerate(assignment, start=1):
        case_id = f"BR{index:03d}"
        terms = [row["source_case_id"].strip(), *row["sensitive_terms"].split("|")]
        available: list[str] = []
        for stage in STAGES:
            raw_path = row[f"{stage}_path"].strip()
            if not raw_path:
                continue
            source = Path(raw_path)
            if not source.is_absolute():
                source = manifest_path.parent / source
            _build_clean_docx(source, output_dir / "documents" / case_id / f"{stage}.docx", terms)
            available.append(stage)
        case_manifest.append({"case_id": case_id, "available_stages": "|".join(available)})
        mapping.append({
            "case_id": case_id,
            "source_case_id": row["source_case_id"].strip(),
            "source_manifest": str(manifest_path.resolve()),
        })

    _write_csv(output_dir / "case_manifest.csv", case_manifest, ("case_id", "available_stages"))
    for reviewer_index, reviewer_code in enumerate(("A", "B"), start=1):
        review_rows = [
            {**{column: "" for column in REVIEW_COLUMNS}, **row, "reviewer_code": reviewer_code, "review_complete": "no"}
            for row in case_manifest
        ]
        random.Random(actual_seed + reviewer_index).shuffle(review_rows)
        _write_csv(output_dir / f"reviewer_{reviewer_code.lower()}.csv", review_rows, REVIEW_COLUMNS)

    private_dir.mkdir(parents=True, exist_ok=True)
    _write_csv(private_dir / "identity_mapping.csv", mapping,
               ("case_id", "source_case_id", "source_manifest"))
    receipt = {
        "schemaVersion": "1.0",
        "createdAt": datetime.now(timezone.utc).isoformat(),
        "caseCount": len(case_manifest),
        "stages": list(STAGES),
        "reviewers": 2,
        "predictionsIncluded": False,
        "textOnlyRebuild": True,
        "identityMappingStoredSeparately": True,
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "package_receipt.json").write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (private_dir / "preparation_receipt.json").write_text(
        json.dumps({**receipt, "shuffleSeed": actual_seed}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return receipt


def _validated_reviews(path: Path) -> dict[str, dict[str, str]]:
    rows = _read_csv(path)
    if not rows:
        raise ValueError(f"标注表为空：{path}")
    missing = set(REVIEW_COLUMNS) - set(rows[0])
    if missing:
        raise ValueError(f"{path.name} 缺少列：{', '.join(sorted(missing))}")
    result: dict[str, dict[str, str]] = {}
    reviewer_codes = set()
    for row in rows:
        case_id = row["case_id"].strip()
        if not case_id or case_id in result:
            raise ValueError(f"{path.name} 含空白或重复 case_id")
        reviewer_codes.add(row["reviewer_code"].strip())
        if row["review_complete"].strip().lower() not in YES_VALUES:
            raise ValueError(f"{path.name} 的 {case_id} 尚未完成")
        if row["overall_risk"].strip().lower() not in REVIEW_LABELS:
            raise ValueError(f"{path.name} 的 {case_id} 风险标签无效")
        if row["scope_change"].strip().lower() not in SCOPE_VALUES:
            raise ValueError(f"{path.name} 的 {case_id} 变更类型无效")
        for column in ("topic_continuity_1_5", "objective_continuity_1_5",
                       "method_continuity_1_5", "confidence_1_5"):
            try:
                value = int(row[column])
            except ValueError as exc:
                raise ValueError(f"{path.name} 的 {case_id}：{column} 必须是 1-5") from exc
            if not 1 <= value <= 5:
                raise ValueError(f"{path.name} 的 {case_id}：{column} 必须是 1-5")
        if not row["evidence_location"].strip() or not row["rationale"].strip():
            raise ValueError(f"{path.name} 的 {case_id} 缺少证据位置或理由")
        row["overall_risk"] = row["overall_risk"].strip().lower()
        result[case_id] = row
    if len(reviewer_codes) != 1:
        raise ValueError(f"{path.name} 必须只包含一名评审人的 reviewer_code")
    return result


def _cohen_kappa(labels_a: list[str], labels_b: list[str]) -> float:
    if not labels_a:
        return 0.0
    observed = sum(a == b for a, b in zip(labels_a, labels_b)) / len(labels_a)
    counts_a, counts_b = Counter(labels_a), Counter(labels_b)
    expected = sum((counts_a[label] / len(labels_a)) * (counts_b[label] / len(labels_b))
                   for label in REVIEW_LABELS)
    return round((observed - expected) / (1 - expected), 4) if expected < 1 else 1.0


def aggregate_reviews(
    reviewer_a_path: Path,
    reviewer_b_path: Path,
    output_dir: Path,
    *,
    adjudication_path: Path | None = None,
) -> dict[str, Any]:
    review_a = _validated_reviews(reviewer_a_path)
    review_b = _validated_reviews(reviewer_b_path)
    if set(review_a) != set(review_b):
        raise ValueError("两位评审人的 case_id 集合不一致")
    reviewer_a_code = next(iter(review_a.values()))["reviewer_code"].strip()
    reviewer_b_code = next(iter(review_b.values()))["reviewer_code"].strip()
    if reviewer_a_code == reviewer_b_code:
        raise ValueError("两份标注表必须来自不同的评审人代码")

    cases = sorted(review_a)
    labels_a = [review_a[case]["overall_risk"] for case in cases]
    labels_b = [review_b[case]["overall_risk"] for case in cases]
    agreements: list[dict[str, str]] = []
    conflicts: list[dict[str, str]] = []
    for case_id in cases:
        a, b = review_a[case_id], review_b[case_id]
        row = {
            "case_id": case_id,
            "reviewer_a_label": a["overall_risk"],
            "reviewer_b_label": b["overall_risk"],
            "reviewer_a_reason": a["rationale"],
            "reviewer_b_reason": b["rationale"],
        }
        if a["overall_risk"] == b["overall_risk"] and a["overall_risk"] in RISK_LABELS:
            agreements.append({"case_id": case_id, "final_risk": a["overall_risk"],
                               "label_source": "reviewer_agreement", "final_reason": "双评审一致"})
        else:
            conflicts.append({**row, "final_risk": "", "final_reason": "", "adjudicator": ""})

    output_dir.mkdir(parents=True, exist_ok=True)
    _write_csv(output_dir / "adjudication.csv", conflicts, (
        "case_id", "reviewer_a_label", "reviewer_b_label", "reviewer_a_reason",
        "reviewer_b_reason", "final_risk", "final_reason", "adjudicator",
    ))
    finalized = list(agreements)
    if adjudication_path:
        decisions = {row["case_id"].strip(): row for row in _read_csv(adjudication_path)}
        if set(decisions) != {row["case_id"] for row in conflicts}:
            raise ValueError("裁决表必须恰好覆盖全部冲突样本")
        for conflict in conflicts:
            decision = decisions[conflict["case_id"]]
            label = decision.get("final_risk", "").strip().lower()
            if label not in RISK_LABELS:
                raise ValueError(f"{conflict['case_id']} 的裁决标签无效")
            if not decision.get("final_reason", "").strip() or not decision.get("adjudicator", "").strip():
                raise ValueError(f"{conflict['case_id']} 缺少裁决理由或裁决人代码")
            finalized.append({
                "case_id": conflict["case_id"],
                "final_risk": label,
                "label_source": "adjudication",
                "final_reason": decision["final_reason"].strip(),
            })

    _write_csv(output_dir / "human_baseline.csv", sorted(finalized, key=lambda item: item["case_id"]),
               ("case_id", "final_risk", "label_source", "final_reason"))
    agreement_count = sum(a == b for a, b in zip(labels_a, labels_b))
    report = {
        "schemaVersion": "1.0",
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "caseCount": len(cases),
        "rawAgreement": round(agreement_count / len(cases), 4),
        "cohenKappa": _cohen_kappa(labels_a, labels_b),
        "conflictCount": len(conflicts),
        "finalizedCount": len(finalized),
        "baselineComplete": len(finalized) == len(cases),
        "reviewerADistribution": dict(Counter(labels_a)),
        "reviewerBDistribution": dict(Counter(labels_b)),
    }
    (output_dir / "agreement_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return report


def write_manifest_template(path: Path) -> None:
    _write_csv(path, [], MANIFEST_COLUMNS)


def main() -> None:
    parser = argparse.ArgumentParser(description="真实匿名样本双教师盲审流程")
    subparsers = parser.add_subparsers(dest="command", required=True)

    manifest_parser = subparsers.add_parser("manifest-template", help="生成私密源清单模板")
    manifest_parser.add_argument("output", type=Path)

    prepare_parser = subparsers.add_parser("prepare", help="制作不含身份映射的盲审包")
    prepare_parser.add_argument("manifest", type=Path)
    prepare_parser.add_argument("output_dir", type=Path)
    prepare_parser.add_argument("private_dir", type=Path)
    prepare_parser.add_argument("--seed", type=int)

    aggregate_parser = subparsers.add_parser("aggregate", help="汇总两位教师的独立标注")
    aggregate_parser.add_argument("reviewer_a", type=Path)
    aggregate_parser.add_argument("reviewer_b", type=Path)
    aggregate_parser.add_argument("output_dir", type=Path)
    aggregate_parser.add_argument("--adjudication", type=Path)

    args = parser.parse_args()
    if args.command == "manifest-template":
        write_manifest_template(args.output)
        print(args.output)
    elif args.command == "prepare":
        print(json.dumps(prepare_package(args.manifest, args.output_dir, args.private_dir,
                                         seed=args.seed), ensure_ascii=False, indent=2))
    else:
        print(json.dumps(aggregate_reviews(args.reviewer_a, args.reviewer_b, args.output_dir,
                                           adjudication_path=args.adjudication),
                         ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
