"""Generate a deterministic, anonymous cross-stage evaluation dataset."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document

ROOT = Path(__file__).resolve().parent
CASES_PATH = ROOT / "cases.json"

TOPICS: dict[str, dict[str, Any]] = {
    "thesis_quality": {"title": "多智能体本科论文质量检查研究", "objective": "构建可解释的论文格式检查、逻辑审查和修改建议生成流程", "method": "采用文档结构解析、确定性规则匹配、任务状态机和多智能体协作", "body": "论文质量检查关注标题层级、章节结构、论证证据、参考文献与排版规范", "keywords": ["论文质检", "多智能体", "规则检查"]},
    "smart_irrigation": {"title": "物联网智能灌溉控制系统研究", "objective": "根据土壤湿度和作物需水量优化农田灌溉策略", "method": "部署湿度传感器并结合模糊控制算法调节阀门流量", "body": "智能灌溉采集土壤墒情、气象数据和作物生长信息以降低农业用水", "keywords": ["智能灌溉", "土壤湿度", "模糊控制"]},
    "campus_energy": {"title": "校园建筑能耗预测与节能研究", "objective": "预测教学楼用电负荷并识别高能耗时段", "method": "融合智能电表、天气和课程安排数据训练时序预测模型", "body": "校园能耗分析覆盖照明、空调、实验设备和分时电价等运行因素", "keywords": ["校园能耗", "负荷预测", "节能"]},
    "cultural_tourism": {"title": "地方文化旅游个性化推荐研究", "objective": "根据游客兴趣和行程约束推荐文化旅游路线", "method": "构建景点知识图谱并结合协同过滤生成路线排序", "body": "文化旅游推荐整合非遗项目、历史景点、游客评价和交通时间", "keywords": ["文化旅游", "推荐系统", "知识图谱"]},
    "inventory_forecast": {"title": "零售库存需求预测与补货优化研究", "objective": "预测商品销量并降低缺货率与库存积压", "method": "利用历史订单、促销和季节特征建立时间序列预测模型", "body": "库存管理关注安全库存、补货周期、需求波动和门店销售数据", "keywords": ["库存预测", "补货优化", "时间序列"]},
    "elderly_health": {"title": "居家老年健康监测与预警研究", "objective": "识别老年人异常生理指标并及时产生健康预警", "method": "采集心率、血氧和活动数据并构建多指标异常检测模型", "body": "居家健康监测关注连续生命体征、跌倒风险、用药提醒和分级预警", "keywords": ["健康监测", "异常检测", "老年人"]},
    "traffic_flow": {"title": "城市道路交通流短时预测研究", "objective": "预测路段拥堵状态并支持信号灯配时优化", "method": "融合道路传感器和路网拓扑构建时空预测模型", "body": "交通流预测分析车速、流量、道路占有率、早晚高峰和相邻路段影响", "keywords": ["交通流", "拥堵预测", "时空模型"]},
    "sentiment_analysis": {"title": "网络评论情感分析方法研究", "objective": "识别网络评论的情感极性和关键意见主题", "method": "清洗评论语料并训练文本表示与情感分类模型", "body": "情感分析处理网络文本、情绪词、否定表达、主题特征和分类标签", "keywords": ["情感分析", "网络评论", "文本分类"]},
    "robot_navigation": {"title": "室内移动机器人自主导航研究", "objective": "实现复杂室内环境中的定位、避障和路径规划", "method": "融合激光雷达建图与启发式搜索算法生成安全路径", "body": "机器人导航涉及同步定位、环境地图、动态障碍、运动控制和路径代价", "keywords": ["机器人导航", "路径规划", "激光雷达"]},
    "water_quality": {"title": "河流水质在线监测与异常识别研究", "objective": "连续监测水质指标并识别污染异常事件", "method": "采集酸碱度、浊度和溶解氧数据并建立异常检测规则", "body": "水质监测分析污染物浓度、传感器漂移、时间变化和异常告警阈值", "keywords": ["水质监测", "污染识别", "传感器"]}
}


def load_cases(path: Path = CASES_PATH) -> list[dict[str, Any]]:
    return json.loads(path.read_text(encoding="utf-8"))


def _add_section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)


def _write_document(path: Path, primary: dict[str, Any], secondary: dict[str, Any] | None,
                    variant: str, stage: str) -> None:
    doc = Document()
    title = primary["title"] if secondary is None else f"{primary['title']}与{secondary['title']}"
    doc.add_heading(title, level=0)
    doc.add_heading("摘要", level=1)
    abstract = primary["body"] * 5 if secondary is None else primary["body"] * 2 + secondary["body"] * 7
    doc.add_paragraph(abstract)
    keywords = primary["keywords"] if secondary is None else primary["keywords"][:1] + secondary["keywords"]
    doc.add_paragraph("关键词：" + "；".join(keywords))
    _add_section(doc, "第一章 绪论", [abstract, primary["body"] * 3])
    objectives = [primary["objective"] * 4]
    methods = [primary["method"] * 4]
    if secondary is not None:
        objectives = [primary["objective"] * 2, secondary["objective"] * 6]
        methods = [primary["method"] * 2, secondary["method"] * 6]
    _add_section(doc, "第二章 研究目标", objectives)
    _add_section(doc, "第三章 研究方法", methods)
    if secondary is not None:
        _add_section(doc, "第四章 新增研究内容", [secondary["body"] * 8])
    evidence = "终稿完成数据整理、对比实验和结果复核。" if stage == "final" else "开题阶段制定数据收集和实验计划。"
    _add_section(doc, "结论", [(secondary or primary)["body"] * 3 + evidence])
    doc.core_properties.subject = f"anonymous-evaluation:{variant}:{stage}"
    doc.core_properties.author = "Anonymous Evaluation Dataset"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def generate_dataset(output_dir: Path, cases_path: Path = CASES_PATH) -> list[dict[str, Any]]:
    cases = load_cases(cases_path)
    generated: list[dict[str, Any]] = []
    for case in cases:
        case_dir = output_dir / case["caseId"]
        source = TOPICS[case["sourceTopic"]]
        target = TOPICS[case["targetTopic"]]
        proposal = case_dir / "proposal.docx"
        final = case_dir / "final.docx"
        _write_document(proposal, source, None, case["variant"], "proposal")
        if case["variant"] == "retained":
            _write_document(final, source, None, case["variant"], "final")
        elif case["variant"] == "partial_shift":
            _write_document(final, source, target, case["variant"], "final")
        else:
            _write_document(final, target, None, case["variant"], "final")
        generated.append({**case, "proposal": proposal, "final": final})
    return generated


if __name__ == "__main__":
    generated = generate_dataset(ROOT / "generated")
    print(f"Generated {len(generated)} anonymous cases in {ROOT / 'generated'}")
