"""DOCX 论文解析模块：提取排版属性与结构信息，供格式/逻辑专家使用。

直接从 WordprocessingML 读取属性（字体、字号、行距、缩进、大纲级别），
避免 python-docx 高层 API 在中文文档中取不到 eastAsia 字体等问题。
"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from docx import Document
from docx.oxml.ns import qn

from tools.logger import get_logger

logger = get_logger("parser")

# 单位换算常量
EMU_PER_CM = 360000.0
TWIP_PER_PT = 20.0
HALF_PT_PER_PT = 2.0
LINE_AUTO_RULE = 240.0

_HEADING_STYLE_RE = re.compile(r"^(?:heading|标题|HEADING)\s*(\d+)$", re.IGNORECASE)

# 段落对齐枚举名到简写的映射（用于样式回退）
_ALIGN_MAP = {"JUSTIFY": "both", "LEFT": "left", "CENTER": "center",
              "RIGHT": "right", "BOTH": "both", "DISTRIBUTE": "both"}
_HEADING_TEXT_RE = re.compile(r"^\s*([0-9]+(?:\.[0-9]+)*|[一二三四五六七八九十]+[、.．]|[（(][一二三四五六七八九十]+[)）])\s*[\s、.．]?")
_CAPTION_RE = re.compile(r"^\s*(图|表|Figure|Fig\.?|Table)\s*([0-9]+[-—－.．][0-9]+|[0-9]+)")
_PUNCT_END = "。；;，,、：:！!？?）)》\"'"


@dataclass
class ParaInfo:
    """段落信息。"""

    index: int
    text: str
    style_id: str = ""
    style_name: str = ""
    level: Optional[int] = None          # 标题层级：1/2/3...
    is_heading: bool = False
    align: Optional[str] = None          # left/center/right/both
    font_cn: Optional[str] = None        # 中文字体（w:eastAsia）
    font_en: Optional[str] = None        # 西文字体（w:ascii）
    size_pt: Optional[float] = None
    bold: Optional[bool] = None
    line_spacing: Optional[float] = None  # 倍数（auto 规则）
    first_indent_chars: Optional[float] = None
    space_before_pt: Optional[float] = None
    space_after_pt: Optional[float] = None
    is_toc_entry: bool = False           # 含 PAGEREF 域（目录项）
    has_image: bool = False
    looks_like_heading: bool = False     # 启发式：疑似手动加粗标题


@dataclass
class Section:
    """按标题切分出的论文章节。"""

    level: int
    title: str
    start_index: int
    end_index: int
    paragraphs: List[ParaInfo] = field(default_factory=list)

    @property
    def text(self) -> str:
        """章节正文文本。"""
        return "\n".join(p.text for p in self.paragraphs if p.text)

    @property
    def char_count(self) -> int:
        """章节中文字符数。"""
        return len(re.findall(r"[一-龥]", self.text))


@dataclass
class DocInfo:
    """论文文档的结构化信息。"""

    path: str = ""
    paragraphs: List[ParaInfo] = field(default_factory=list)
    sections: List[Section] = field(default_factory=list)
    page: Dict[str, Optional[float]] = field(default_factory=dict)
    toc: Dict[str, object] = field(default_factory=dict)
    headers: List[str] = field(default_factory=list)
    footers: List[str] = field(default_factory=list)
    has_page_number: bool = False
    references: List[str] = field(default_factory=list)
    ref_heading_index: Optional[int] = None
    captions: List[Dict[str, object]] = field(default_factory=list)
    table_count: int = 0
    image_count: int = 0
    full_text: str = ""
    stats: Dict[str, int] = field(default_factory=dict)
    body_size_pt: Optional[float] = None
    body_items: List[Tuple[str, int]] = field(default_factory=list)
    body_start_index: Optional[int] = None  # 正文起始段落（封面/声明/摘要/目录之后）
    abstract_cn: str = ""
    abstract_en: str = ""
    keywords: List[str] = field(default_factory=list)
    conclusion_text: str = ""

    @property
    def body_paragraphs(self) -> List[ParaInfo]:
        """正文段落（排除标题与空段落）。"""
        return [p for p in self.paragraphs if p.text.strip() and not p.is_heading]


def _xml_val(node, attr: str = "w:val") -> Optional[str]:
    """读取 XML 节点的属性值，自动补全 w: 命名空间。

    Args:
        node: XML 节点。
        attr: 属性名。

    Returns:
        属性值字符串，不存在时返回 None。
    """
    if node is None:
        return None
    name = attr if ":" in attr else f"w:{attr}"
    return node.get(qn(name))


def _to_float(value: Optional[str]) -> Optional[float]:
    """字符串转浮点数，失败返回 None。"""
    try:
        return float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None


def _align_of(pPr) -> Optional[str]:
    """读取段落对齐方式。"""
    node = pPr.find(qn("w:jc")) if pPr is not None else None
    return _xml_val(node)


def _spacing_of(pPr) -> Dict[str, Optional[float]]:
    """读取段间距与行距。

    Returns:
        {line: 倍数, before: pt, after: pt}。
    """
    result: Dict[str, Optional[float]] = {"line": None, "before": None, "after": None}
    if pPr is None:
        return result
    node = pPr.find(qn("w:spacing"))
    if node is None:
        return result
    line = _to_float(_xml_val(node, "w:line"))
    rule = _xml_val(node, "w:lineRule")
    if line is not None:
        if rule == "auto" or rule is None:
            result["line"] = round(line / LINE_AUTO_RULE, 2)
        else:  # exact / atLeast：单位为二十分之一磅
            result["line"] = round(line / TWIP_PER_PT, 2)
    before = _to_float(_xml_val(node, "w:before"))
    after = _to_float(_xml_val(node, "w:after"))
    result["before"] = round(before / TWIP_PER_PT, 2) if before is not None else None
    result["after"] = round(after / TWIP_PER_PT, 2) if after is not None else None
    return result


def _indent_chars(pPr) -> Optional[float]:
    """读取首行缩进（优先 firstLineChars，否则按 firstLine 磅值估算字符数）。"""
    if pPr is None:
        return None
    node = pPr.find(qn("w:ind"))
    if node is None:
        return None
    chars = _to_float(_xml_val(node, "w:firstLineChars"))
    if chars is not None:
        return round(chars / 100.0, 2)
    twips = _to_float(_xml_val(node, "w:firstLine"))
    if twips is not None:
        return round(twips / TWIP_PER_PT / 12.0, 2)  # 粗略按 12pt 字宽折算
    return None


def _font_of(rPr) -> Dict[str, Optional[object]]:
    """读取字体属性（中/西文字体、字号、加粗）。"""
    result: Dict[str, Optional[object]] = {"cn": None, "en": None, "size": None, "bold": None}
    if rPr is None:
        return result
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is not None:
        result["cn"] = _xml_val(fonts, "w:eastAsia")
        result["en"] = _xml_val(fonts, "w:ascii") or _xml_val(fonts, "w:hAnsi")
    sz = rPr.find(qn("w:sz"))
    size = _to_float(_xml_val(sz))
    result["size"] = round(size / HALF_PT_PER_PT, 1) if size is not None else None
    bold_node = rPr.find(qn("w:b"))
    if bold_node is not None:
        flag = _xml_val(bold_node)
        result["bold"] = flag not in ("0", "false", "off")
    return result


def _detect_level(paragraph, style_id: str, style_name: str, pPr) -> Optional[int]:
    """判定段落大纲级别。

    优先读取 w:outlineLvl，其次解析样式名（Heading N / 标题 N）。
    """
    if pPr is not None:
        node = pPr.find(qn("w:outlineLvl"))
        raw = _xml_val(node)
        if raw is not None and raw.isdigit():
            level = int(raw) + 1
            if 1 <= level <= 9:
                return level
    for name in (style_id, style_name):
        match = _HEADING_STYLE_RE.match((name or "").strip())
        if match:
            level = int(match.group(1))
            if 1 <= level <= 9:
                return level
    return None


def _page_settings(sections: Sequence) -> Dict[str, Optional[float]]:
    """取论文主体页面设置（按节取众数，避免封面单独页边距干扰）。

    Args:
        sections: python-docx 的节对象序列。

    Returns:
        {width_cm, height_cm, top_cm, bottom_cm, left_cm, right_cm}。
    """
    sizes: Counter = Counter()
    margins: Counter = Counter()
    for sec in sections:
        try:
            if sec.page_width and sec.page_height:
                sizes[(round(sec.page_width / EMU_PER_CM, 1),
                       round(sec.page_height / EMU_PER_CM, 1))] += 1
            values = (sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin)
            if all(v is not None for v in values):
                margins[tuple(round(v / EMU_PER_CM, 1) for v in values)] += 1
        except Exception:  # noqa: BLE001 - 单个节异常不影响整体
            continue
    if not sizes:
        return {}
    width, height = sizes.most_common(1)[0][0]
    top = bottom = left = right = None
    if margins:
        top, bottom, left, right = margins.most_common(1)[0][0]
    return {"width_cm": width, "height_cm": height,
            "top_cm": top, "bottom_cm": bottom, "left_cm": left, "right_cm": right}


def _header_footer_texts(sections: Sequence) -> Tuple[List[str], List[str]]:
    """收集全部节的页眉与页脚文字。

    Args:
        sections: python-docx 的节对象序列。

    Returns:
        (页眉文字列表, 页脚文字列表)，已去重。
    """
    headers: List[str] = []
    footers: List[str] = []
    for sec in sections:
        for container, bucket in ((sec.header, headers), (sec.footer, footers)):
            try:
                if container is None or container.is_linked_to_previous:
                    continue
                for para in container.paragraphs:
                    text = (para.text or "").strip()
                    if text and text not in bucket:
                        bucket.append(text)
            except Exception:  # noqa: BLE001
                continue
    return headers, footers


def parse_docx(path: str | Path) -> DocInfo:
    """解析 docx 论文文件。

    Args:
        path: docx 文件路径。

    Returns:
        DocInfo 结构化信息对象。

    Raises:
        FileNotFoundError: 文件不存在。
        ValueError: 文件不是合法 docx 或不含段落。
    """
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"论文文件不存在：{file_path}")

    document = Document(str(file_path))
    info = DocInfo(path=str(file_path))

    # ---------- 页面设置 ----------
    try:
        info.page = _page_settings(document.sections)
        info.headers, info.footers = _header_footer_texts(document.sections)
        footer_xml = "".join(sec.footer._element.xml for sec in document.sections
                             if sec.footer is not None)
        info.has_page_number = ("PAGE" in footer_xml) or bool(info.footers)
    except Exception as exc:  # noqa: BLE001 - 页面设置缺失不影响主体解析
        logger.debug("页面设置解析失败：%s", exc)
        info.page = {}
        info.has_page_number = False
        info.headers, info.footers = [], []

    # ---------- 正文元素顺序（段落 / 表格交错） ----------
    para_seq = 0
    tbl_seq = 0
    for child in document.element.body.iterchildren():
        tag = str(child.tag).split("}")[-1]
        if tag == "p":
            info.body_items.append(("p", para_seq))
            para_seq += 1
        elif tag == "tbl":
            info.body_items.append(("tbl", tbl_seq))
            tbl_seq += 1

    # ---------- 段落 ----------
    toc_field = False
    toc_entries = 0
    for idx, para in enumerate(document.paragraphs):
        xml_para = para._p
        pPr = xml_para.find(qn("w:pPr"))
        try:
            style_id = para.style.style_id or ""
            style_name = para.style.name or ""
        except Exception:  # noqa: BLE001 - 样式缺失时退回空串
            style_id, style_name = "", ""

        # 直接格式优先，缺失时回退到样式定义
        direct = _font_of(pPr.find(qn("w:rPr")) if pPr is not None else None)
        size_pt = direct["size"]
        font_cn = direct["cn"]
        font_en = direct["en"]
        bold = direct["bold"]
        if size_pt is None or font_cn is None:
            try:
                style_font = para.style.font
                if size_pt is None and style_font.size is not None:
                    size_pt = round(style_font.size.pt, 1)
                if font_cn is None and style_font.name:
                    font_cn = style_font.name
                    font_en = font_en or style_font.name
            except Exception:  # noqa: BLE001
                pass

        spacing = _spacing_of(pPr)
        level = _detect_level(para, style_id, style_name, pPr)
        text = para.text or ""

        # 直接格式缺失时回退到样式定义（行距、首行缩进）
        if spacing["line"] is None:
            try:
                ls = para.paragraph_format.line_spacing
                if isinstance(ls, (int, float)):
                    spacing["line"] = round(float(ls), 2)
            except Exception:  # noqa: BLE001
                pass
        first_indent = _indent_chars(pPr)
        if first_indent is None:
            try:
                fi = para.paragraph_format.first_line_indent
                if fi is not None and hasattr(fi, "pt"):
                    first_indent = round(fi.pt / (size_pt or 12.0), 2)
            except Exception:  # noqa: BLE001
                pass

        # 目录域与页码域
        instr = " ".join(
            node.text or "" for node in xml_para.iter(qn("w:instrText")) if node.text
        )
        if instr:
            upper = instr.upper()
            if "TOC" in upper:
                toc_field = True
            if "PAGEREF" in upper:
                toc_entries += 1

        has_image = bool(xml_para.findall(f".//{qn('w:drawing')}")) or bool(
            xml_para.findall(f".//{qn('w:pict')}")
        )

        # 直接格式缺失时回退到样式定义（对齐方式）
        align = _align_of(pPr)
        if align is None:
            try:
                effective = para.paragraph_format.alignment
                if effective is not None:
                    align = _ALIGN_MAP.get(str(getattr(effective, "name", effective)).upper())
            except Exception:  # noqa: BLE001
                pass

        item = ParaInfo(
            index=idx,
            text=text,
            style_id=style_id,
            style_name=style_name,
            level=level,
            is_heading=level is not None and bool(text.strip()),
            align=align,
            font_cn=font_cn,
            font_en=font_en,
            size_pt=size_pt,
            bold=bold,
            line_spacing=spacing["line"],
            first_indent_chars=first_indent,
            space_before_pt=spacing["before"],
            space_after_pt=spacing["after"],
            is_toc_entry="PAGEREF" in instr.upper(),
            has_image=has_image,
        )
        info.paragraphs.append(item)

        match = _CAPTION_RE.match(text.strip())
        if match and len(text.strip()) < 80:
            info.captions.append(
                {
                    "kind": "图" if match.group(1) in ("图", "Figure", "Fig.", "Fig") else "表",
                    "number": match.group(2),
                    "text": text.strip(),
                    "index": idx,
                }
            )

    if not info.paragraphs:
        raise ValueError("文档未包含任何段落，可能不是有效的 docx 论文")

    # 众数正文字号（用于启发式识别「伪标题」）
    sizes = [p.size_pt for p in info.paragraphs if p.size_pt and p.text.strip() and not p.is_heading]
    if sizes:
        info.body_size_pt = max(set(sizes), key=sizes.count)

    for para in info.paragraphs:
        para.looks_like_heading = _looks_like_heading(para, info.body_size_pt)

    # ---------- 章节切分 ----------
    info.sections = _build_sections(info.paragraphs)

    # ---------- 目录 ----------
    heading_levels = [p.level for p in info.paragraphs if p.is_heading and p.level]
    info.toc = {
        "has_field": toc_field,
        "entries": toc_entries,
        "has_entries": toc_entries >= 3,
        "heading_levels": sorted(set(heading_levels)),
        "max_level": max(heading_levels) if heading_levels else 0,
    }

    # ---------- 参考文献 ----------
    info.ref_heading_index = _find_reference_heading(info.paragraphs)
    info.references = _extract_references(info.paragraphs, info.ref_heading_index)

    # ---------- 统计 ----------
    try:
        info.table_count = len(document.tables)
        info.image_count = len(document.inline_shapes)
    except Exception:  # noqa: BLE001
        info.table_count = info.image_count = 0

    info.body_start_index = _find_body_start(info.paragraphs)
    info.abstract_cn, info.abstract_en, info.keywords = _extract_abstract(info)
    info.conclusion_text = _extract_conclusion(info.sections)

    info.full_text = "\n".join(p.text for p in info.paragraphs if p.text)
    cn_chars = len(re.findall(r"[一-龥]", info.full_text))
    en_words = len(re.findall(r"[A-Za-z]+", info.full_text))
    digits = len(re.findall(r"\d", info.full_text))
    info.stats = {
        "chars_cn": cn_chars,
        "words_en": en_words,
        "digits": digits,
        "total_chars": len(re.sub(r"\s", "", info.full_text)),
        "paragraphs": len([p for p in info.paragraphs if p.text.strip()]),
        "headings": len([p for p in info.paragraphs if p.is_heading]),
    }
    return info


_BODY_START_RE = re.compile(
    r"^\s*(第\s*[0-9一二三四五六七八九十百]+\s*[章节篇]|绪论|引言|前言"
    r"|[0-9]+\s*[、.．]\s*\S)")


def _find_body_start(paragraphs: List[ParaInfo]) -> Optional[int]:
    """定位正文起始段落（跳过封面、声明、摘要、目录）。

    Args:
        paragraphs: 全部段落。

    Returns:
        正文首段下标；未识别到章节标题时返回 None。
    """
    for para in paragraphs:
        if para.is_heading and _BODY_START_RE.match(para.text.strip()):
            return para.index
    return None


def _text_until_next_heading(paragraphs: List[ParaInfo], start_index: Optional[int]) -> str:
    """提取指定标题之后、下一个同级或更高级标题之前的正文。

    Args:
        paragraphs: 全部段落。
        start_index: 起始标题段落下标，None 时返回空串。

    Returns:
        拼接后的正文文本。
    """
    if start_index is None:
        return ""
    level: Optional[int] = None
    for para in paragraphs:
        if para.index == start_index:
            level = para.level
            break
    texts: List[str] = []
    for para in paragraphs:
        if para.index <= start_index:
            continue
        if para.is_heading and level and para.level and para.level <= level:
            break
        if para.text.strip():
            texts.append(para.text.strip())
    return "\n".join(texts)


def _expand_from_following(paragraphs: List[ParaInfo], start_index: Optional[int],
                           current: str, min_len: int = 60, window: int = 12) -> str:
    """摘要文本过短时向后补足（模板中常紧邻说明性段落）。

    Args:
        paragraphs: 全部段落。
        start_index: 摘要标题下标。
        current: 已提取的摘要文本。
        min_len: 触发补足的最短长度。
        window: 向后取段落的数量。

    Returns:
        补足后的文本。
    """
    if len(current) >= min_len or start_index is None:
        return current
    extra = [p.text.strip() for p in paragraphs
             if start_index < p.index <= start_index + window and p.text.strip()]
    return (current + "\n" + "\n".join(extra)).strip()


def _extract_abstract(info: DocInfo) -> Tuple[str, str, List[str]]:
    """提取中英文摘要与关键词。

    Args:
        info: 文档信息。

    Returns:
        (中文摘要, 英文摘要, 关键词列表)。
    """
    cn_idx: Optional[int] = None
    en_idx: Optional[int] = None
    for para in info.paragraphs:
        if not para.is_heading:
            continue
        text = para.text.strip()
        if en_idx is None and re.search(r"(英文摘要|ABSTRACT|Abstract)", text):
            en_idx = para.index
        elif cn_idx is None and re.search(r"摘\s*要", text):
            cn_idx = para.index

    # 未使用标题样式时，按纯文本回退查找「摘要 / Abstract」行
    if cn_idx is None or en_idx is None:
        for para in info.paragraphs:
            text = para.text.strip()
            if len(text) > 20:
                continue
            if en_idx is None and re.fullmatch(r"(英文摘要|ABSTRACT|Abstract)", text):
                en_idx = para.index
            elif cn_idx is None and re.fullmatch(r"摘\s*要", text):
                cn_idx = para.index

    cn_text = _expand_from_following(info.paragraphs, cn_idx,
                                     _text_until_next_heading(info.paragraphs, cn_idx))
    en_text = _expand_from_following(info.paragraphs, en_idx,
                                     _text_until_next_heading(info.paragraphs, en_idx))

    keywords: List[str] = []
    match = re.search(r"关键词\s*[:：]\s*(.+)", cn_text)
    if match:
        keywords = [k.strip() for k in re.split(r"[；;，,、\s]+", match.group(1)) if k.strip()]
        cn_text = cn_text[: match.start()]
    if not keywords:
        for para in info.paragraphs:
            m = re.match(r"\s*关键词[:：]\s*(.+)", para.text.strip())
            if m:
                keywords = [k.strip() for k in re.split(r"[；;，,、\s]+", m.group(1)) if k.strip()]
                break
    return cn_text.strip(), en_text.strip(), keywords


def _extract_conclusion(sections: List[Section]) -> str:
    """提取结论章节正文。

    Args:
        sections: 章节列表。

    Returns:
        结论章节文本，未找到时返回空串。
    """
    for section in sections:
        if re.search(r"(结论|结语|结束语|总结与展望)", section.title):
            return section.text
    return ""


def _looks_like_heading(para: ParaInfo, body_size: Optional[float]) -> bool:
    """启发式判断是否为手动加粗的「伪标题」。

    Args:
        para: 段落信息。
        body_size: 正文众数字号。

    Returns:
        疑似标题返回 True。
    """
    text = para.text.strip()
    if not text or para.is_heading:
        return False
    if len(text) > 50 or text[-1] in _PUNCT_END:
        return False
    if not para.bold:
        return False
    if body_size and para.size_pt and para.size_pt < body_size:
        return False
    return bool(_HEADING_TEXT_RE.match(text))


def _find_reference_heading(paragraphs: List[ParaInfo]) -> Optional[int]:
    """定位参考文献标题段落下标。"""
    for para in paragraphs:
        if not para.is_heading:
            continue
        text = para.text.strip()
        if re.search(r"(参考文献|References|REFERENCES)", text):
            return para.index
    return None


def _extract_references(paragraphs: List[ParaInfo], start_index: Optional[int]) -> List[str]:
    """提取参考文献条目文本。"""
    if start_index is None:
        return []
    start_level = None
    for para in paragraphs:
        if para.index == start_index:
            start_level = para.level
            break
    items: List[str] = []
    for para in paragraphs:
        if para.index <= start_index:
            continue
        if para.is_heading and start_level and para.level and para.level <= start_level:
            break
        text = para.text.strip()
        if len(text) >= 8:
            items.append(text)
    return items


def _build_sections(paragraphs: List[ParaInfo]) -> List[Section]:
    """按标题层级切分章节。

    Args:
        paragraphs: 全部段落。

    Returns:
        章节列表（不含层级为 None 的段落）。
    """
    sections: List[Section] = []
    current: Optional[Section] = None
    for para in paragraphs:
        if para.is_heading and para.level:
            if current is not None:
                current.end_index = para.index - 1
            current = Section(
                level=para.level,
                title=para.text.strip(),
                start_index=para.index,
                end_index=len(paragraphs) - 1,
            )
            sections.append(current)
        elif current is not None:
            current.paragraphs.append(para)
    return sections
